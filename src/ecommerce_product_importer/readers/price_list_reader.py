from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook


@dataclass
class PriceRow:
    model: str
    color: str
    size: str
    net_price: float
    recommended_price: float


def read_price_list(file_path: Path) -> list[PriceRow]:
    """
    Read a supplier price list from Excel file.
    """
    if not file_path.exists():
        raise FileNotFoundError(f"Price list not found: {file_path}")

    workbook = load_workbook(file_path, data_only=True)
    worksheet = workbook.active

    rows: list[PriceRow] = []

    # Skip the header row
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue

        rows.append(
            PriceRow(
                model=str(row[0]),
                color=str(row[1]),
                size=str(row[2]),
                net_price=float(row[3]),
                recommended_price=float(row[4]),
            )
        )

    return rows
