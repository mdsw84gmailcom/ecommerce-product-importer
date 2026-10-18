from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


REQUIRED_COLUMNS = {
    "Product name (sv)",
    "Product number",
    "Brand",
    "Category (sv)",
    "Description (sv)",
    "Product images",
    "Main image",
    "Rek. ut",
    "ÅF-pris",
    "Tags",
    "Color (sv)",
    "Color code",
    "Sku",
    "Size code",
    "Size name",
    "Ean",
    "Available from",
}


@dataclass(frozen=True)
class DJProductRow:
    """One sellable D&J article/SKU normalized from the Excel export."""

    supplier: str
    product_id: str
    product_name: str
    brand: str
    category: str
    description: str
    image_urls: tuple[str, ...]
    main_image: str | None
    recommended_price: float | None
    dealer_price: float | None
    gender: str | None
    fit: str | None
    material_technique: str | None
    season: str | None
    tags: tuple[str, ...]
    packaging_comment: str | None
    presentation_date: date | None
    color: str
    color_code: str
    sku: str
    size_code: str
    size_name: str
    ean: str | None
    available_from: date | None


def _text(value: Any, default: str = "") -> str:
    if value is None:
        return default
    return str(value).strip()


def _optional_text(value: Any) -> str | None:
    value_text = _text(value)
    return value_text or None


def _number(value: Any) -> float | None:
    if value in (None, ""):
        return None
    return float(value)


def _date(value: Any) -> date | None:
    if value in (None, ""):
        return None

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    value_text = _text(value)
    for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(value_text, fmt).date()
        except ValueError:
            continue

    raise ValueError(f"Unsupported D&J date value: {value!r}")


def _split(value: Any, separator: str = ",") -> tuple[str, ...]:
    if value in (None, ""):
        return ()

    return tuple(
        item.strip()
        for item in str(value).split(separator)
        if item.strip()
    )


def read_dj_export(file_path: Path) -> list[DJProductRow]:
    """Read a D&J Excel data export and return normalized article rows."""
    if not file_path.exists():
        raise FileNotFoundError(f"D&J export not found: {file_path}")

    if file_path.suffix.lower() != ".xlsx":
        raise ValueError("The D&J export must be an .xlsx file.")

    workbook = load_workbook(file_path, data_only=True, read_only=True)
    worksheet = workbook.active

    header_values = next(
        worksheet.iter_rows(min_row=1, max_row=1, values_only=True)
    )
    headers = [_text(value) for value in header_values]
    column_index = {name: index for index, name in enumerate(headers)}

    missing_columns = REQUIRED_COLUMNS - set(column_index)
    if missing_columns:
        missing = ", ".join(sorted(missing_columns))
        raise ValueError(f"Missing required D&J columns: {missing}")

    def value(row: tuple[Any, ...], column: str) -> Any:
        index = column_index.get(column)
        if index is None or index >= len(row):
            return None
        return row[index]

    rows: list[DJProductRow] = []

    for row in worksheet.iter_rows(min_row=2, values_only=True):
        product_id = _text(value(row, "Product number"))
        sku = _text(value(row, "Sku"))

        if not product_id and not sku:
            continue

        if not product_id:
            raise ValueError(
                f"D&J row with SKU {sku!r} is missing Product number"
            )

        if not sku:
            raise ValueError(f"D&J product {product_id!r} is missing SKU")

        rows.append(
            DJProductRow(
                supplier="D&J",
                product_id=product_id,
                product_name=_text(value(row, "Product name (sv)")),
                brand=_text(value(row, "Brand")),
                category=_text(value(row, "Category (sv)")),
                description=_text(value(row, "Description (sv)")),
                image_urls=_split(value(row, "Product images")),
                main_image=_optional_text(value(row, "Main image")),
                recommended_price=_number(value(row, "Rek. ut")),
                dealer_price=_number(value(row, "ÅF-pris")),
                gender=_optional_text(value(row, "Gender (sv)")),
                fit=_optional_text(value(row, "Fit (sv)")),
                material_technique=_optional_text(
                    value(row, "Material technique (sv)")
                ),
                season=_optional_text(value(row, "Season")),
                tags=_split(value(row, "Tags"), separator=";"),
                packaging_comment=_optional_text(
                    value(row, "Packaging Comment (sv)")
                ),
                presentation_date=_date(value(row, "Presentation date")),
                color=_text(value(row, "Color (sv)")),
                color_code=_text(value(row, "Color code")),
                sku=sku,
                size_code=_text(value(row, "Size code")),
                size_name=_text(value(row, "Size name")),
                ean=_optional_text(value(row, "Ean")),
                available_from=_date(value(row, "Available from")),
            )
        )

    return rows
